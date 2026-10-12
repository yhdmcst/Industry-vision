"""Historical draft builder, not the assembler for the current competition bundle.

Embedded schedules and deadlines are unverified legacy development text. Do not
use this script to replace the required PPT or to package the current directory.
The current submission instructions live in the outer 00_提交说明 directory.
"""

import hashlib
import json
import math
from pathlib import Path
import re
import unicodedata
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
import yaml

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "artifacts/submission"
INK, TEAL, GOLD, MUTED = "202B30", "007C72", "C98D21", "657278"
BG = "F6F7F8"
BENCHMARK_NAMES = ("benchmark.json", "benchmark-4b.json",
                   "benchmark-initial-failure.json", "benchmark-v1-failures.json")


def color(value):
    return RGBColor.from_string(value)


def text(slide, value, x, y, w, h, size=20, bold=False, fill=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(str(value).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(fill)
        paragraph.line_spacing = 1.15
        paragraph.space_after = Pt(6)
    return box


def block(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background()
    return shape


def load_benchmarks() -> tuple[dict, dict[str, bytes]]:
    benchmarks = {}
    records = {}
    for name in BENCHMARK_NAMES:
        path = ROOT / "artifacts" / name
        if not path.exists():
            continue
        records[name] = path.read_bytes()
        data = json.loads(records[name])
        if not isinstance(data, dict) or not isinstance(data.get("cases"), list):
            raise ValueError(f"{name}: cases must be a list")
        for index, case in enumerate(data["cases"], 1):
            try:
                for label in [case["scenario_ground_truth"], case["rule_analysis"]["drift_type"],
                              case["analysis"]["drift_type"]]:
                    if not isinstance(label, str) or not label.strip():
                        raise ValueError("scenario and classification labels must be nonempty strings")
                for label, value in [("elapsed_ms", case["inference"]["elapsed_ms"]),
                                     ("vision p50", case["vision_latency_ms"]["p50"])]:
                    if (isinstance(value, bool) or not isinstance(value, (int, float))
                            or not math.isfinite(value) or value < 0):
                        raise ValueError(f"{label} must be finite and nonnegative")
            except (KeyError, TypeError, ValueError) as error:
                raise ValueError(f"{name}, case {index}: {error}") from error
        benchmarks[name] = data
    return benchmarks, records


def submission_files() -> list[Path]:
    selected = []
    for directory in ["src", "scripts", "tests"]:
        selected.extend(path for path in (ROOT / directory).rglob("*")
                        if path.suffix.lower() in {".py", ".ps1"})
    for directory in ["frontend/src", "frontend/public"]:
        selected.extend(path for path in (ROOT / directory).rglob("*")
                        if path.suffix.lower() in {".ts", ".tsx", ".css", ".svg"})
    names = [
        "frontend/package.json", "frontend/package-lock.json", "frontend/index.html",
        "frontend/tsconfig.json", "frontend/vite.config.ts", "docs/ui-design.md", "docs/datasets.md", "docs/inspection-results.md", "docs/evaluation.md", "docs/reliability.md", "docs/acquisition.md",
        "README.md", "SUMMARY.md", "ROADMAP.md", "run.py",
        "docs/composition.md", "docs/public-data.md", "docs/visual-detector.md", "docs/vision-validation.md", "docs/vision-acceptance.md", "docs/review-evaluation.md", "docs/repository.md", "requirements-vision.txt",
        "requirements.txt", "requirements-model.txt", "requirements-lock.txt",
        "config/system_config.yaml",
        "docs/team_profile.yaml", "docs/submission.md",
        "docs/deployment.md", "docs/third_party.md", "docs/originality.md", "docs/verification.md",
        "docs/sources/spark-official-readme.md", "docs/sources/spark-gguf-files.json",
        "docs/sources/spark-4b-gguf-files.json",
        "models/manifest.json", "models/manifest-4b.json", "models/research-qwen3-4b.json", "models/Modelfile",
        "docs/soak.md", "docs/operations.md",
        "models/Modelfile-4b", "models/runtime-macos.json", "models/runtime-platforms.json", "artifacts/ui-test.json",
        "artifacts/thinking-comparison.json", "artifacts/thinking-comparison-2048-failures.json",
        "artifacts/device-camera.json",
    ]
    names.extend(f"artifacts/{name}" for name in BENCHMARK_NAMES)
    names.extend(f"artifacts/screenshots/{name}.png" for name in ["desktop", "mobile", "model-review"])
    selected.extend(ROOT / name for name in names)
    excluded = {"__pycache__", "cache", "caches", "logs", "log", "private", "original-proposal"}
    root = ROOT.resolve()
    return sorted({path for path in selected if path.is_file()
                   and path.resolve().is_relative_to(root)
                   and not any(part.lower() in excluded or part.startswith(".")
                               for part in path.relative_to(ROOT).parts)})


def check_structure(pres: Presentation) -> dict:
    checks = {"geometry_errors": [], "text_bounds_errors": [], "text_overlap_errors": []}

    def check_frame(frame, width, height, slide_index, name):
        if not frame.text.strip():
            return
        available = (width - frame.margin_left - frame.margin_right) / 12700
        required = (frame.margin_top + frame.margin_bottom) / 12700
        for index, paragraph in enumerate(frame.paragraphs):
            size = max([run.font.size.pt for run in paragraph.runs if run.font.size]
                       + [paragraph.font.size.pt if paragraph.font.size else 18])
            # Conservative glyph widths and explicit line spacing, not a rendered-font measurement.
            lines, used = 1, 0.0
            for char in paragraph.text:
                advance = size * (1 if unicodedata.east_asian_width(char) in {"W", "F"} else .6)
                if used and used + advance > available:
                    lines, used = lines + 1, 0.0
                used += advance
            spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else 1.2
            required += lines * size * spacing
            required += paragraph.space_before.pt if paragraph.space_before else 0
            if index < len(frame.paragraphs) - 1:
                required += paragraph.space_after.pt if paragraph.space_after else 0
        if required > height / 12700 + .5 or available <= 0:
            checks["text_bounds_errors"].append({
                "slide": slide_index, "shape": name, "estimated_height_pt": round(required, 2),
                "height_pt": round(height / 12700, 2),
            })

    for index, slide in enumerate(pres.slides, 1):
        content = []
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0 or shape.left + shape.width > pres.slide_width
                    or shape.top + shape.height > pres.slide_height):
                checks["geometry_errors"].append({"slide": index, "shape": shape.name, "issue": "outside slide"})
            if shape.has_text_frame and shape.text.strip():
                check_frame(shape.text_frame, shape.width, shape.height, index, shape.name)
                content.append(shape)
            if shape.has_table:
                content.append(shape)
                for row_index, row in enumerate(shape.table.rows):
                    for column_index, cell in enumerate(row.cells):
                        check_frame(cell.text_frame, shape.table.columns[column_index].width,
                                    row.height, index, f"{shape.name}/r{row_index}/c{column_index}")
        for position, first in enumerate(content):
            for second in content[position + 1:]:
                width = min(first.left + first.width, second.left + second.width) - max(first.left, second.left)
                height = min(first.top + first.height, second.top + second.height) - max(first.top, second.top)
                if width > Inches(.01) and height > Inches(.01):
                    checks["text_overlap_errors"].append({
                        "slide": index, "shapes": [first.name, second.name],
                    })
    return checks


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    team = yaml.safe_load((ROOT / "docs/team_profile.yaml").read_text(encoding="utf-8"))
    benchmarks, benchmark_records = load_benchmarks()
    benchmark_source = next((name for name in ["benchmark-4b.json", "benchmark.json"]
                             if benchmarks.get(name, {}).get("cases")), None)
    primary = benchmarks.get(benchmark_source, {})
    cases = primary.get("cases", [])
    model_name = primary.get("model_manifest", {}).get("official_model", "模型实测待完成")
    manifest = primary.get("model_manifest", {})
    pres = Presentation()
    pres.slide_width, pres.slide_height = Inches(13.333), Inches(7.5)
    pres.core_properties.title = "Spark Guard 工业质检复核台 | 初赛草稿"
    pres.core_properties.subject = "Measured local prototype, not factory validation"
    pres.core_properties.author = team["team_name"]

    def new_slide(title, note, dark=False):
        slide = pres.slides.add_slide(pres.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(INK if dark else BG)
        text(slide, title, .65, .55, 12, .8, 32, True, "FFFFFF" if dark else INK)
        text(slide, "SPARK GUARD / 初赛项目介绍草稿", .65, 7.03, 10, .25, 10, fill=MUTED if not dark else "BFCBCD")
        text(slide, f"{len(pres.slides):02d}", 12.05, 7.0, .55, .3, 12, fill=MUTED)
        slide.notes_slide.notes_text_frame.text = note
        return slide

    slide = new_slide("Spark Guard · 工业质检复核台",
                      "这是可运行工程原型，不是已完成产线验证的产品。负责人和团队信息必须由真实成员填写。", True)
    text(slide, "先核验检测条件\n再讨论产品质量", .7, 1.7, 6.6, 1.8, 38, True, "FFFFFF")
    text(slide, f"{model_name} / 本机文本证据复核\n{team['team_name']}\n团队资料与原创性声明待核验", .7, 4.3, 6.6, 1.65, 20, fill="D3DDDE")
    block(slide, 8.0, 1.65, 4.6, 4.75, TEAL)
    text(slide, "旁路观测\n竞争解释\n人工验证\n证据留存", 8.4, 2.05, 3.8, 3.9, 28, True, "FFFFFF")

    slide = new_slide("团队与资源", "所有履历、公司架构和合作资源均待提供；不得编造。")
    text(slide, f"负责人：{team['leader']['name']}\n职责：{team['leader']['role']}\n履历：{team['leader']['biography']}",
         .7, 1.7, 6.6, 3.8, 23)
    block(slide, 8.1, 1.65, 4.4, 4.75, "E5ECEB")
    text(slide, "真实信息待人工核验", 8.4, 2, 3.8, .6, 24, True, TEAL)
    members = f"{len(team['members'])} 项，未核验" if team["members"] else "未提供"
    partners = f"{len(team['partners'])} 项，未核验" if team["partners"] else "未提供"
    text(slide, f"团队架构：{team['organization']}\n核心成员：{members}\n合作伙伴：{partners}\n联系方式：{team['contact']}",
         8.4, 2.9, 3.8, 3.1, 17)

    slide = new_slide("NG 告警增加，不一定是产品变差",
                      "用户价值是假设，尚无工厂客户验证与ROI测量。目标用户为视觉集成商和现场维护人员。")
    for x, title, body, fill in [
        (.7, "成像条件变化", "失焦 / 污染\n光照 / 曝光\n工装 / 摆位", TEAL),
        (4.95, "检测告警增加", "相似度变化\n稳定性变化\nNG 代理上升", GOLD),
        (9.15, "产品变化待查", "人工抽检\n批次核查\n独立缺陷标签", INK),
    ]:
        block(slide, x, 1.9, 3.5, .65, fill)
        text(slide, title, x + .15, 2.03, 3.2, .45, 22, True, "FFFFFF")
        text(slide, body, x + .1, 2.95, 3.3, 2.2, 23)
    text(slide, "价值目标：缩小排查范围，而不是替代质量负责人作出最终判断。", .7, 5.95, 11.8, .65, 22, True)

    slide = new_slide("完整的旁路复核流程", "所有设备控制均在范围之外。环境传感器和真实外部质检系统尚未接入。")
    steps = [
        ("01 采集", "正常基线 / 当前帧", TEAL), ("02 测量", "质量 / 配准 / 一致性", INK),
        ("03 复核", "Spark + 操作员记录", GOLD), ("04 留证", "SQLite / JSON / 历史", TEAL),
    ]
    for index, (title, body, fill) in enumerate(steps):
        x = .7 + index * 3.15
        block(slide, x, 2.0, 2.8, .85, fill)
        text(slide, title, x + .18, 2.2, 2.45, .5, 24, True, "FFFFFF")
        text(slide, body, x, 3.25, 2.8, 1.2, 20)
        if index < 3:
            text(slide, ">", x + 2.88, 2.2, .25, .5, 24, True, MUTED)
    text(slide, "快路径：视觉指标与规则观测\n慢路径：单任务异步模型复核，绑定测量快照", .7, 5.25, 11.8, 1.2, 22)

    slide = new_slide("模型与基准记录中的部署", "文本证据复核，不宣称原生多模态。本页读取基准记录，不是重新测量或硬件核验。")
    inference = cases[0]["inference"] if cases else {}
    runtime = inference.get("runtime", {})
    loaded_models = primary.get("loaded_model_evidence", {}).get("models", [])
    loaded = next((item for item in loaded_models
                   if item.get("name") == inference.get("model")), {})
    backend = "Ollama" if inference.get("backend") == "ollama" else inference.get("backend", "运行时待测")
    memory = primary.get("memory_total_bytes", 0)
    ram = f"{memory / 1024**3:.1f} GiB RAM" if memory else "内存待测"
    text(slide, f"{model_name}\n{runtime.get('format', '格式待测')} / {runtime.get('quantization_level', '量化待测')}\n"
         f"{backend} {runtime.get('runtime_version', '版本待测')} / {loaded.get('context_length', '待测')} tokens\n"
         f"{primary.get('platform', '平台待测')} / {ram}",
         .7, 1.65, 6.4, 3.4, 24)
    block(slide, 8.0, 1.65, 4.5, 4.8, "E5ECEB")
    text(slide, "真实性链路", 8.35, 2.0, 3.8, .55, 25, True, TEAL)
    source_status = "仓库版本：见基准清单" if manifest.get("revision") else "模型来源清单待提供"
    sha_status = ("历史基准权重摘要与清单一致" if runtime.get("weight_sha256") == manifest.get("sha256")
                  and manifest.get("sha256") else "基准权重摘要核验待完成")
    ps_status = "基准含 /api/ps 运行记录" if loaded else "/api/ps 记录待提供"
    text(slide, f"{source_status}\n{sha_status}\n{ps_status}\n原始响应与计时：见基准", 8.35, 2.9, 3.8, 3, 20)
    sha = manifest.get("sha256")
    text(slide, f"SHA-256: {sha[:24]}..." if sha else "SHA-256: 待核验", .7, 5.65, 7.0, .6, 14, fill=MUTED)
    text(slide, f"记录来源：artifacts/{benchmark_source}" if benchmark_source else "无模型实测记录；不报告部署成绩",
         .7, 6.35, 11.8, .4, 14, fill=MUTED)

    slide = new_slide("模型负责复核，不负责设备执行", "模型贡献是竞争解释与可执行核验计划；不可替代性尚未被实验验证。")
    for x, title, body in [
        (.7, "规则观测", "定位指标变化\n快速触发复核\n可审计的阈值"),
        (4.9, "Spark 复核", "结合操作员记录\n提出备选原因\n组织区分性验证"),
        (9.1, "人工决策", "检查实际图像\n采集独立证据\n决定维护与质量处置"),
    ]:
        text(slide, title, x, 1.9, 3.5, .6, 25, True, TEAL)
        text(slide, body, x, 3, 3.5, 2.4, 22)
    text(slide, "模型可能误判。结构正确、语言流畅，不等于根因正确。", .7, 6.0, 11.8, .6, 23, True, GOLD)

    slide = new_slide("运行界面与证据导出", "历史界面：2026-10-05 合成样本模型复核；当前版本使用 React。")
    screenshot = ROOT / "artifacts/screenshots/model-review.png"
    if screenshot.exists():
        with Image.open(screenshot) as image:
            width, height = image.size
        picture_width = min(8.0, 4.9 * width / height)
        slide.shapes.add_picture(str(screenshot), Inches(.7), Inches(1.55),
                                 width=Inches(picture_width), height=Inches(picture_width * height / width))
    else:
        text(slide, "截图待采集；此处不展示虚构界面", .7, 2.5, 8, 1, 24, fill=MUTED)
    ui_path = ROOT / "artifacts/ui-test.json"
    ui = json.loads(ui_path.read_text(encoding="utf-8")) if ui_path.exists() else {}
    ui_results = ui.get("results", [])
    ui_recorded = (isinstance(ui_results, list) and "page_errors" in ui and not ui["page_errors"]
                   and "desktop loaded and single-frame sampling succeeded" in ui_results
                   and "mobile 390px layout has no horizontal page overflow" in ui_results)
    ui_status = "桌面采样 / 手机尺寸：有记录" if ui_recorded else "桌面 / 手机验证记录待补全"
    text(slide, f"基线与当前帧\n规则与模型分开显示\n冲突提示 / 人工复核\n原始输出 / 历史 / JSON\n{ui_status}\n"
         "记录来源：ui-test.json\n后端测试：见 tests/\n本页不代表本次重跑测试",
         9.0, 1.85, 3.5, 4.55, 18)

    single_calls = len({case["scenario_ground_truth"] for case in cases}) == len(cases)
    sampling_note = ("每场景一次模型调用；无置信区间；调试样例非盲测。" if single_calls
                     else "逐项展示模型调用记录；无置信区间；调试样例非盲测。")
    for start in range(0, max(len(cases), 1), 6):
        page_cases = cases[start:start + 6]
        slide = new_slide("实测结果：视觉与模型分别计时",
                          f"来源：{benchmark_source or '无基准'}。{sampling_note}不作为准确率或工厂验证。")
        if page_cases:
            synthetic = all(case.get("snapshot", {}).get("synthetic") is True for case in page_cases)
            headers = ["合成场景" if synthetic else "记录场景", "规则类别", "模型类别", "模型耗时 / s"]
            table_shape = slide.shapes.add_table(len(page_cases) + 1, 4, Inches(.7), Inches(1.9),
                                                Inches(11.8), Inches(3.4))
            table = table_shape.table
            for index, header in enumerate(headers):
                table.cell(0, index).text = header
            for row_index, case in enumerate(page_cases, 1):
                values = [case["scenario_ground_truth"], case["rule_analysis"]["drift_type"],
                          case["analysis"]["drift_type"], f"{case['inference']['elapsed_ms']/1000:.2f}"]
                for index, value in enumerate(values):
                    table.cell(row_index, index).text = value
            for row_index, row in enumerate(table.rows):
                for cell in row.cells:
                    cell.margin_top = cell.margin_bottom = Inches(.05)
                    cell.text_frame.word_wrap = True
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = "Microsoft YaHei"
                        paragraph.font.size = Pt(18)
                        paragraph.line_spacing = 1.15
                        paragraph.font.color.rgb = color("FFFFFF" if row_index == 0 else INK)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = color(TEAL if row_index == 0 else "FFFFFF")
        else:
            text(slide, "模型实测待完成；不报告耗时或准确率", .7, 2.5, 11.8, 1, 24, fill=MUTED)
        vision = [case["vision_latency_ms"]["p50"] for case in page_cases]
        text(slide, f"本页各场景视觉 p50：{min(vision):.1f}–{max(vision):.1f} ms" if vision else "视觉测量待完成",
             .7, 5.6, 11.8, .55, 21, True, TEAL)
        text(slide, f"来源：{benchmark_source or '待测'}；逐项保留结论（含误判），早期记录见 ZIP。",
             .7, 6.18, 11.8, .4, 15, fill=MUTED)
        text(slide, sampling_note if cases else "未提供模型实测记录。",
             .7, 6.65, 11.8, .3, 13, fill=MUTED)

    slide = new_slide("验证设计与消融", "不能把规则阈值提示和合成训练式调试当成大模型不可替代性证明。")
    for y, name, content in [
        (1.8, "输入消融", "相同指标：无操作员记录 vs 有真实维护记录"),
        (3.0, "模型对照", "规则 / 1.7B / 4B / 同档备选模型，固定采集集"),
        (4.2, "盲测评价", "类别、数字忠实度、验证可执行性、错误建议"),
        (5.4, "真实设备", "独立产品标签、断网测试、冷 / 热延迟与功耗"),
    ]:
        block(slide, .7, y, 2.1, .65, TEAL)
        text(slide, name, .88, y + .11, 1.75, .45, 21, True, "FFFFFF")
        text(slide, content, 3.2, y + .08, 9.1, .65, 21)

    slide = new_slide("工程与安全边界", "本应用没有公网登录和租户隔离，不应对外直接暴露。")
    text(slide, "仅本机 / 权重核验\n无 Mock / 结构校验\n旧基线结果丢弃\n证据冲突 / 人工复核\n低信号拒绝校准\n原始响应与错误留存", .7, 1.7, 5.4, 4.2, 23)
    block(slide, 7.3, 1.65, 5.2, 4.85, "E5ECEB")
    text(slide, "尚未验证", 7.65, 2.0, 4.5, .6, 27, True, GOLD)
    text(slide, "工厂场景与收益\n有效取景 / 实体故障\n断网与功耗实测\n跨设备部署\n模型替代 / 消融增益", 7.65, 2.95, 4.5, 3.2, 21)

    slide = new_slide("提交冲刺：先完成真实证据",
                      "按用户提供文本较早截止2026-10-11准备；另有2026-11-10 17:00，必须向组委会确认。")
    for x, date, task in [
        (.7, "10.05–07", "真摄像头\n真实质检结果\n盲测采集"),
        (3.85, "10.08–09", "对照与消融\n断网验证\n团队与授权"),
        (7.0, "10.10", "3–5分钟视频\nPPT人工复核\n另一设备复现"),
        (10.15, "10.11前", "确认截止\n提交ZIP\n回执备份"),
    ]:
        text(slide, date, x, 2, 2.5, .7, 28, True, TEAL)
        text(slide, task, x, 3.25, 2.5, 2.5, 20)
    text(slide, "赛程文本有 10月11日 / 11月10日17:00 冲突，请组委会确认。", .7, 6.15, 11.8, .5, 17, fill=GOLD)

    slide = new_slide("原创性与真实性声明草稿",
                      "待负责人审核、署名或盖章后提交；此页不能表示已经签署。")
    text(slide, "承诺范围", .7, 1.7, 3, .7, 28, True, TEAL)
    text(slide, "如实呈现本团队完成的工作；不将模型和上游框架宣称为独立原创。\n区分合成、软件扰动与真实采集，不伪造指标、履历、合作资源或现场证据。\n核验所有权和授权，知识产权纠纷按赛事规则由参赛者承担。",
         .7, 2.8, 11.8, 2.9, 23)
    block(slide, .7, 6.05, 11.8, .6, "E5ECEB")
    text(slide, "团队 / 负责人 / 日期 / 签名或盖章：待填写与签署", .9, 6.15, 11.4, .4, 18, True, GOLD)

    ppt = OUTPUT / "SparkGuard-初赛项目介绍-草稿.pptx"
    pres.save(ppt)
    if not (ROOT / "requirements-lock.txt").is_file():
        raise RuntimeError("Resolve and review requirements-lock.txt before packaging")
    checks = check_structure(Presentation(ppt))
    team_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", team["team_name"]).strip(" .") or "待填写团队名称"
    bundle = OUTPUT / f"{team_name}+SparkGuard-初赛草稿.zip"
    selected = [path for path in submission_files()
                if path.parent != ROOT / "artifacts" or path.name not in BENCHMARK_NAMES
                or path.name in benchmark_records]
    packaging_skipped = any(checks.values())
    if not packaging_skipped:
        with zipfile.ZipFile(bundle, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            for path in selected:
                name = path.relative_to(ROOT).as_posix()
                if path.parent == ROOT / "artifacts" and path.name in benchmark_records:
                    archive.writestr(name, benchmark_records[path.name])
                else:
                    archive.write(path, name)
            archive.write(ppt, "artifacts/submission/" + ppt.name)
    required_team = [team.get("team_name"), team.get("contact"), team.get("organization"),
                     *(team.get("leader", {}).get(key) for key in ["name", "role", "biography"])]
    missing_team = (any(not value or not str(value).strip() for value in required_team)
                    or any(marker in str(team) for marker in ["待填写", "待提供", "TBD", "TODO"]))
    report = {"slides": len(pres.slides), **checks,
              "pptx": ppt.name, "zip": None if packaging_skipped else bundle.name,
              "zip_bytes": 0 if packaging_skipped else bundle.stat().st_size,
              "packaging_skipped": packaging_skipped,
              "draft": True, "missing_team_information": missing_team,
              "benchmark_source": benchmark_source,
              "benchmark_sha256": hashlib.sha256(benchmark_records[benchmark_source]).hexdigest()
              if benchmark_source else None,
              "zip_files": [] if packaging_skipped else [path.relative_to(ROOT).as_posix() for path in selected]
              + ["artifacts/submission/" + ppt.name],
              "visual_review_performed": False,
              "note": "Structural geometry and conservative text estimates only; no rendered visual review."}
    (OUTPUT / "build-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True))
    return 1 if any(checks.values()) else 0


if __name__ == "__main__":
    raise SystemExit(main())
