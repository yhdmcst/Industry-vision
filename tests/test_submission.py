"""Independent structural QA; these tests do not render or inspect slide images."""

import copy
import hashlib
import json
from pathlib import Path
import shutil
import zipfile
from xml.etree import ElementTree

from pptx import Presentation
from pptx.util import Inches
import pytest
import yaml

from scripts import build_submission as builder

PROJECT = Path(__file__).resolve().parents[1]


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    for name in [
        "README.md", "SUMMARY.md", "ROADMAP.md", "run.py",
        "docs/composition.md", "docs/public-data.md", "docs/visual-detector.md", "docs/vision-validation.md", "docs/vision-acceptance.md", "docs/review-evaluation.md", "docs/repository.md", "requirements-vision.txt",
        "requirements.txt", "requirements-model.txt", "requirements-lock.txt", "docs/team_profile.yaml", "docs/operations.md",
        "models/manifest.json", "models/Modelfile", "models/runtime-macos.json", "models/runtime-platforms.json", "models/research-qwen3-4b.json", "docs/soak.md", "config/system_config.yaml", "docs/datasets.md", "docs/inspection-results.md", "docs/evaluation.md", "docs/reliability.md", "docs/acquisition.md",
        "frontend/package.json", "frontend/package-lock.json", "frontend/src/App.tsx", "artifacts/benchmark.json", "artifacts/ui-test.json",
    ]:
        destination = tmp_path / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(PROJECT / name, destination)
    for name in ["src/example.py", "scripts/example.ps1", "tests/example.py"]:
        destination = tmp_path / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text("# source fixture\n", encoding="utf-8")
    monkeypatch.setattr(builder, "ROOT", tmp_path)
    monkeypatch.setattr(builder, "OUTPUT", tmp_path / "artifacts/submission")
    return tmp_path


def read_outputs(root):
    output = root / "artifacts/submission"
    report = json.loads((output / "build-report.json").read_text(encoding="utf-8"))
    presentation = Presentation(output / report["pptx"])
    return report, presentation, output / report["zip"]


def deck_text(presentation):
    return "\n".join(shape.text for slide in presentation.slides
                     for shape in slide.shapes if shape.has_text_frame)


def table_rows(presentation):
    return [
        [cell.text for cell in row.cells]
        for slide in presentation.slides for shape in slide.shapes if shape.has_table
        for row in list(shape.table.rows)[1:]
    ]


def expected_rows(benchmark):
    return [
        [case["scenario_ground_truth"], case["rule_analysis"]["drift_type"],
         case["analysis"]["drift_type"], f"{case['inference']['elapsed_ms'] / 1000:.2f}"]
        for case in benchmark["cases"]
    ]


def test_actual_benchmark_survives_pptx_roundtrip(workspace):
    benchmark_path = workspace / "artifacts/benchmark.json"
    lock_before = (workspace / "requirements-lock.txt").read_bytes()
    before = benchmark_path.read_bytes()
    benchmark = json.loads(before)
    assert builder.main() == 0
    report, presentation, bundle = read_outputs(workspace)
    assert table_rows(presentation) == expected_rows(benchmark)
    content = deck_text(presentation)
    assert "28 项后端测试" not in content
    assert "待填写" in content
    assert report["missing_team_information"] is True
    assert report["draft"] is True
    assert not report["geometry_errors"]
    assert not report["text_bounds_errors"]
    assert not report["text_overlap_errors"]
    assert report["visual_review_performed"] is False
    assert report["benchmark_sha256"] == hashlib.sha256(before).hexdigest()
    assert benchmark_path.read_bytes() == before
    assert (workspace / "requirements-lock.txt").read_bytes() == lock_before
    with zipfile.ZipFile(bundle) as archive:
        assert "docs/datasets.md" in archive.namelist()
        assert "docs/inspection-results.md" in archive.namelist()
        assert "docs/evaluation.md" in archive.namelist()
        assert "docs/reliability.md" in archive.namelist()
        assert "models/research-qwen3-4b.json" in archive.namelist()
        assert "docs/vision-acceptance.md" in archive.namelist()
        assert "docs/repository.md" in archive.namelist()
        assert "docs/soak.md" in archive.namelist()
        assert "docs/operations.md" in archive.namelist()
        assert "docs/acquisition.md" in archive.namelist()
        assert "docs/composition.md" in archive.namelist()
        assert "docs/public-data.md" in archive.namelist()
        assert "docs/visual-detector.md" in archive.namelist()
        assert "requirements-vision.txt" in archive.namelist()
        assert "docs/review-evaluation.md" in archive.namelist()
        assert "models/runtime-macos.json" in archive.namelist()
        assert archive.testzip() is None
        assert archive.read("artifacts/benchmark.json") == before
        assert "requirements-model.txt" in archive.namelist()
        assert archive.read("requirements-lock.txt") == (workspace / "requirements-lock.txt").read_bytes()


def test_bundle_excludes_weights_caches_secrets_and_private_logs(workspace):
    excluded = [
        "frontend/node_modules/private.ts", "frontend/dist/private.js", "frontend/.env", "models/weights.gguf", "models/private.json", "models/ollama/blobs/weights",
        "scripts/__pycache__/example.pyc", "scripts/.cache/private.py",
        "scripts/logs/private.py", "tests/private.log", "config/secrets.yaml",
        ".streamlit/secrets.toml", "docs/private.md", "docs/original-proposal/README.md",
        "artifacts/ui-evidence.json", "artifacts/private.json", "artifacts/benchmark.db",
        ".venv/private.py", "data/db/private.db", "runtime/private.log", ".env",
    ]
    for name in excluded:
        destination = workspace / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"PRIVATE_SENTINEL")
    assert builder.main() == 0
    _, _, bundle = read_outputs(workspace)
    with zipfile.ZipFile(bundle) as archive:
        names = set(archive.namelist())
        assert not names.intersection(excluded)
        assert {"models/manifest.json", "models/Modelfile", "src/example.py",
                "scripts/example.ps1", "tests/example.py", "frontend/src/App.tsx",
                "frontend/package.json", "frontend/package-lock.json"} <= names
        assert all(b"PRIVATE_SENTINEL" not in archive.read(name) for name in names)


def test_missing_benchmark_never_claims_measured_deployment(workspace):
    (workspace / "artifacts/benchmark.json").unlink()
    assert builder.main() == 0
    report, presentation, _ = read_outputs(workspace)
    content = deck_text(presentation)
    assert not table_rows(presentation)
    assert "模型实测待完成" in content
    assert "Ollama 0.35.1" not in content
    assert "RTX 5060" not in content
    assert report["benchmark_source"] is None


def test_empty_4b_benchmark_does_not_replace_measured_17b(workspace):
    (workspace / "artifacts/benchmark-4b.json").write_text('{"cases": []}', encoding="utf-8")
    assert builder.main() == 0
    report, presentation, _ = read_outputs(workspace)
    benchmark = json.loads((workspace / "artifacts/benchmark.json").read_text(encoding="utf-8"))
    assert report["benchmark_source"] == "benchmark.json"
    assert table_rows(presentation) == expected_rows(benchmark)


def test_4b_deployment_and_rows_use_selected_evidence(workspace):
    benchmark = json.loads((workspace / "artifacts/benchmark.json").read_text(encoding="utf-8"))
    benchmark["model_manifest"]["official_model"] = "Spark-X2.5-4B"
    benchmark["model_manifest"]["model_name"] = "spark-drift:4b-q4"
    for model in benchmark["loaded_model_evidence"]["models"]:
        model["name"] = "spark-drift:4b-q4"
        model["context_length"] = 2048
    for case in benchmark["cases"]:
        case["inference"]["model"] = "spark-drift:4b-q4"
        case["inference"]["runtime"]["runtime_version"] = "test-runtime"
        case["inference"]["elapsed_ms"] += 1000
    (workspace / "artifacts/benchmark-4b.json").write_text(
        json.dumps(benchmark, ensure_ascii=False), encoding="utf-8")
    assert builder.main() == 0
    report, presentation, _ = read_outputs(workspace)
    content = deck_text(presentation)
    assert report["benchmark_source"] == "benchmark-4b.json"
    assert table_rows(presentation) == expected_rows(benchmark)
    assert "Spark-X2.5-4B" in content
    assert "test-runtime" in content
    assert "2048 tokens" in content
    assert "4096 tokens" not in content


@pytest.mark.parametrize("elapsed", [-1, float("inf"), float("nan")])
def test_invalid_timing_is_not_published(workspace, elapsed):
    path = workspace / "artifacts/benchmark.json"
    benchmark = json.loads(path.read_text(encoding="utf-8"))
    benchmark["cases"][0]["inference"]["elapsed_ms"] = elapsed
    path.write_text(json.dumps(benchmark), encoding="utf-8")
    with pytest.raises(ValueError, match="elapsed_ms"):
        builder.main()
    assert not list((workspace / "artifacts/submission").glob("*.zip"))


def test_benchmark_rows_paginate_without_losing_results(workspace):
    path = workspace / "artifacts/benchmark.json"
    benchmark = json.loads(path.read_text(encoding="utf-8"))
    benchmark["cases"] = copy.deepcopy(benchmark["cases"] * 3)
    path.write_text(json.dumps(benchmark), encoding="utf-8")
    assert builder.main() == 0
    report, presentation, _ = read_outputs(workspace)
    assert table_rows(presentation) == expected_rows(benchmark)
    assert all(len(shape.table.rows) <= 7 for slide in presentation.slides
               for shape in slide.shapes if shape.has_table)
    assert not report["text_bounds_errors"]
    assert "一次调用 / 场景" not in deck_text(presentation)


@pytest.mark.parametrize("missing", [False, True])
def test_failed_or_missing_ui_report_does_not_claim_workflow_pass(workspace, missing):
    path = workspace / "artifacts/ui-test.json"
    if missing:
        path.unlink()
    else:
        path.write_text('{"results": [], "page_errors": ["failure"]}', encoding="utf-8")
    assert builder.main() == 0
    _, presentation, _ = read_outputs(workspace)
    assert "桌面与手机流程通过" not in deck_text(presentation)


def test_structural_checker_detects_bounds_overflow_and_text_overlap():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    builder.text(slide, "Outside", 13, 1, 1, 1)
    builder.text(slide, "Too much text\n" * 8, 1, 1, 2, .3, 24)
    builder.text(slide, "Overlap", 1.5, 1, 2, .3, 20)
    checks = builder.check_structure(presentation)
    assert checks["geometry_errors"]
    assert checks["text_bounds_errors"]
    assert checks["text_overlap_errors"]


def test_overlong_team_content_blocks_packaging(workspace):
    path = workspace / "docs/team_profile.yaml"
    team = yaml.safe_load(path.read_text(encoding="utf-8"))
    team["leader"]["biography"] = "Long biography " * 300
    path.write_text(yaml.safe_dump(team, allow_unicode=True), encoding="utf-8")
    assert builder.main() == 1
    output = workspace / "artifacts/submission"
    report = json.loads((output / "build-report.json").read_text(encoding="utf-8"))
    assert report["text_bounds_errors"]
    assert report["packaging_skipped"] is True
    assert report["zip"] is None
    assert not list(output.glob("*.zip"))


def test_project_snapshot_with_actual_screenshot_and_early_records(tmp_path, monkeypatch):
    for source in builder.submission_files():
        destination = tmp_path / source.relative_to(PROJECT)
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, destination)
    monkeypatch.setattr(builder, "ROOT", tmp_path)
    monkeypatch.setattr(builder, "OUTPUT", tmp_path / "artifacts/submission")
    assert builder.main() == 0
    report, presentation, bundle = read_outputs(tmp_path)
    benchmark = json.loads((tmp_path / "artifacts" / report["benchmark_source"]).read_bytes())
    assert table_rows(presentation) == expected_rows(benchmark)
    assert report["visual_review_performed"] is False
    assert "待填写与签署" in deck_text(presentation)
    assert not any(report[key] for key in
                   ["geometry_errors", "text_bounds_errors", "text_overlap_errors"])
    with zipfile.ZipFile(bundle) as archive:
        assert archive.testzip() is None
        for name in builder.BENCHMARK_NAMES:
            source = tmp_path / "artifacts" / name
            if source.exists():
                assert archive.read(f"artifacts/{name}") == source.read_bytes()
        assert all(not name.lower().endswith((".gguf", ".pyc", ".db", ".log"))
                   for name in archive.namelist())
    with zipfile.ZipFile(tmp_path / "artifacts/submission" / report["pptx"]) as archive:
        assert archive.testzip() is None
        media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
        assert len(media) == 1
        assert archive.read(media[0]) == (tmp_path / "artifacts/screenshots/model-review.png").read_bytes()
        namespaces = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for name in archive.namelist():
            if name.endswith(".xml"):
                document = ElementTree.fromstring(archive.read(name))
                if name.startswith("ppt/slides/slide"):
                    for transform in document.findall(".//a:xfrm", namespaces):
                        offset = transform.find("a:off", namespaces)
                        extent = transform.find("a:ext", namespaces)
                        if offset is not None and extent is not None:
                            x, y = int(offset.attrib["x"]), int(offset.attrib["y"])
                            width, height = int(extent.attrib["cx"]), int(extent.attrib["cy"])
                            assert x >= 0 and y >= 0 and width > 0 and height > 0
                            assert x + width <= presentation.slide_width
                            assert y + height <= presentation.slide_height
